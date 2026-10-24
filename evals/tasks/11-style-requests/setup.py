import json, random, sys
import openpyxl

work, truth = sys.argv[1], sys.argv[2]
random.seed(11)
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Orders"
ws.append(["Order ID", "Region", "Rep", "Units", "Price", "Revenue", "Status"])
late = []
for i in range(2, 32):
    status = random.choice(["Open", "Shipped", "Shipped", "Late"])
    ws.append([f"SO-{1000 + i}", random.choice(["North", "South", "East", "West"]), random.choice(["Ann", "Bo", "Cy", "Di"]),
               random.randint(1, 40), round(random.uniform(5, 250), 2), f"=D{i}*E{i}", status])
    if status == "Late":
        late.append(i)
ws.append(["Total", None, None, "=SUM(D2:D31)", None, "=SUM(F2:F31)", None])
wb.save(f"{work}/orders.xlsx")
revenue = {r: ws.cell(r, 4).value * ws.cell(r, 5).value for r in range(2, 32)}
json.dump({"late": late, "rows": list(range(2, 32)), "total_row": 32, "revenue": revenue, "f2": revenue[2]}, open(truth, "w"))
