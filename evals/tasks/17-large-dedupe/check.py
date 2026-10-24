import openpyxl
from evalkit import Checks, args, close, mentions, stored, values

work, truth, answer = args()
path = f"{work}/orders_2026.xlsx"
c = Checks()
if not c.guard("opens", lambda: (openpyxl.load_workbook(path, read_only=True) is not None, "")):
    c.report()
    raise SystemExit
wb = openpyxl.load_workbook(path, read_only=True)
ids, units = [], 0
for r in wb["Data"].iter_rows(min_row=2, values_only=True):
    if r[0] is None:
        continue
    ids.append(r[0])
    units += r[5] or 0
c.add("each order once", len(ids) == len(set(ids)), f"{len(ids)} rows, {len(set(ids))} distinct")
c.add("no other order removed", len(set(ids)) == truth["rows"] and units == truth["units"], f"{len(set(ids))} ids, {units} units")
full = openpyxl.load_workbook(path)
f5, f15 = full["Summary"]["B5"].value, full["Summary"]["B15"].value
c.add("the Summary is still formulas", all(isinstance(f, str) and f.startswith("=") for f in (f5, f15)), (f5, f15))
s, _ = values(path, "Summary!B5:B15")
c.add("Summary March revenue", close(s[5][0], truth["march"], 1e-6), f"{s[5][0]} want {truth['march']}")
c.add("Summary full-year total", close(s[15][0], truth["total"], 1e-6), f"{s[15][0]} want {truth['total']}")
c.add("answer gives the March revenue", mentions(answer, truth["march"], 1e-4), f"want {truth['march']:,.2f}")
c.add("answer gives the full-year total", mentions(answer, truth["total"], 1e-4), f"want {truth['total']:,.2f}")
d = full["Data"]
c.add("the header keeps its style", d["A1"].font.bold and (d["A1"].fill.fgColor.rgb or "").endswith("1F4E79"), (d["A1"].font.bold, d["A1"].fill.fgColor.rgb))
c.add("revenue keeps its number format", d["H2"].number_format == "#,##0.00", d["H2"].number_format)
c.add("the header stays frozen and filtered", d.freeze_panes == "A2" and bool(d.auto_filter.ref), (d.freeze_panes, d.auto_filter.ref))
v = stored(path, "Summary", "B5")
c.add("Summary March stored", close(v, truth["march"], 1e-6), v, kind="stored")
c.report()
