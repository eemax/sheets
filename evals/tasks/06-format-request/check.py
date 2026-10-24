import openpyxl
from evalkit import Checks, args

work, truth, answer = args()
wb = openpyxl.load_workbook(f"{work}/budget.xlsx")
ws = wb["Budget"]
c = Checks()
head = [ws.cell(1, col) for col in range(1, 5)]
c.add("header bold", all(cell.font.b for cell in head))
def blue(cell):
    rgb = (cell.fill.fgColor.rgb or "") if cell.fill and cell.fill.fill_type == "solid" else ""
    if not isinstance(rgb, str) or len(rgb) < 6:
        return False
    r, g, b = (int(rgb[-6:][i:i + 2], 16) for i in (0, 2, 4))
    return b >= r and b >= g and b > 150 and r + g + b > 450
c.add("header light blue", all(blue(cell) for cell in head), [cell.fill.fgColor.rgb for cell in head])
fmts = [ws.cell(r, col).number_format for r in range(2, 5) for col in range(2, 5)]
c.add("dollars without decimals", all("$" in f and ".0" not in f and ".#" not in f for f in fmts), sorted(set(fmts)))
c.add("header frozen", ws.freeze_panes == "A2", ws.freeze_panes)
c.add("data intact", ws["D3"].value == "=B3-C3" and ws["B2"].value == 1200000, (ws["B2"].value, ws["D3"].value))
c.report()
