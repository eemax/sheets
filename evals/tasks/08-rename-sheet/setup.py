import json, sys
import openpyxl
from openpyxl.workbook.defined_name import DefinedName

work, truth = sys.argv[1], sys.argv[2]
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Data"
for row in [("Region", "Units"), ("North", 5), ("South", 7), ("East", 3), ("West", 9)]:
    ws.append(row)
ws["D1"] = 2.5
wb.defined_names["Rate"] = DefinedName("Rate", attr_text="Data!$D$1")
s = wb.create_sheet("Summary")
s["A1"], s["B1"] = "All units", "=SUM(Data!B2:B5)"
s["A2"], s["B2"] = "North doubled", "=Data!B2*2"
s["A3"], s["B3"] = "South at rate", "=Data!B3*Rate"
wb.save(f"{work}/book.xlsx")
json.dump({"B1": 24, "B2": 10, "B3": 17.5}, open(truth, "w"))
