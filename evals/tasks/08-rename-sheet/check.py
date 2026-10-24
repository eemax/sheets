import openpyxl
from evalkit import Checks, args, close, stored, values

work, truth, answer = args()
path = f"{work}/book.xlsx"
c = Checks()
wb = openpyxl.load_workbook(path)
c.add("sheet renamed", "Sales 2026" in wb.sheetnames and "Data" not in wb.sheetnames, wb.sheetnames)
s, _ = values(path, "Summary!A1:B3")
for i, cell in enumerate(["B1", "B2", "B3"], 1):
    c.add(f"Summary!{cell} still computes", close(s[i][1], truth[cell]), s[i][1])
    c.add(f"Summary!{cell} stored", close(stored(path, "Summary", cell), truth[cell]), stored(path, "Summary", cell), kind="stored")
rate = wb.defined_names.get("Rate")
c.add("the name Rate follows", rate is not None and "Sales 2026" in rate.attr_text, rate.attr_text if rate else None)
c.report()
