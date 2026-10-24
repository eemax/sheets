import re
import zipfile
import openpyxl
from openpyxl.utils import range_boundaries
from evalkit import Checks, args, close, formula, stored, values

work, truth, answer = args()
path = f"{work}/dashboard.xlsx"
c = Checks()
if not c.guard("opens", lambda: (openpyxl.load_workbook(path) is not None, "")):
    c.report()
    raise SystemExit
rows, _ = values(path, "Dashboard!A1:G30")
by = {v[0]: (r, v) for r, v in rows.items() if v and v[0]}
c.add("North Q4 is 1,520", "North" in by and close(by["North"][1][4], 1520), by.get("North"))
c.add("South Q4 is 980", "South" in by and close(by["South"][1][4], 980), by.get("South"))
c.add("North's total follows", "North" in by and close(by["North"][1][5], truth["north"]), by.get("North"))
nordics = by.get("Nordics")
c.add("Nordics has its quarters", nordics and [round(float(x)) for x in nordics[1][1:5]] == [410, 455, 470, 512], nordics)
c.add("Nordics' total is a formula", nordics and close(nordics[1][5], truth["nordics"]) and formula(path, "Dashboard", f"F{nordics[0]}"),
      nordics and (nordics[1][5], formula(path, "Dashboard", f"F{nordics[0]}")))
t = openpyxl.load_workbook(path)["Dashboard"].tables.get("Sales")
c.add("Nordics is in the table", t is not None and nordics and range_boundaries(t.ref)[3] >= nordics[0], t and t.ref)

# What the workbook held stays: parts openpyxl drops or rewrites are read
# from the package itself.
z = zipfile.ZipFile(path)
names = z.namelist()
sheet = z.read("xl/worksheets/sheet1.xml").decode()
drawings = "".join(z.read(n).decode() for n in names if re.match(r"xl/drawings/drawing\d+\.xml$", n))
parts = "".join(z.read(n).decode(errors="ignore") for n in names if n.endswith(".xml") or n.endswith(".vml"))
wbx = z.read("xl/workbook.xml").decode()
c.add("the chart is kept", any(n.startswith("xl/charts/chart") for n in names) and "Sales by quarter" in parts)
c.add("the logo is kept", any(n.startswith("xl/media/") for n in names) and re.search(r"<(\w+:)?pic>", drawings) is not None)
c.add("the text box is kept", "ERP export" in drawings)
c.add("the sparklines are kept", sheet.count("sparkline>") >= 4 * 2, sheet.count("sparkline>") // 2)
c.add("the data bars keep their Excel 2010 look", "x14:dataBar" in sheet)
c.add("the dropdown is kept", "On track,At risk,Behind" in sheet.replace("&quot;", '"').replace('"', ""))
c.add("the note is kept", "Baltics" in parts)
c.add("the link is kept", "hyperlink" in sheet and "example.com/finance" in "".join(z.read(n).decode() for n in names if n.endswith(".rels")))
c.add("the panes stay frozen", 'state="frozen"' in sheet)
c.add("the print header is kept", re.search(r"<oddHeader>[^<]*Regional dashboard", sheet) is not None)
c.add("the print area is kept", "_xlnm.Print_Area" in wbx)
if "North" in by:
    v = stored(path, "Dashboard", f"F{by['North'][0]}")
    c.add("North total stored", close(v, truth["north"]), v, kind="stored")
c.report()
