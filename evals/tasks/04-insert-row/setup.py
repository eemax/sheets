# A small model: data, a total, and a summary sheet reading single cells.
import json, sys
import openpyxl

work, truth = sys.argv[1], sys.argv[2]
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Data"
for row in [("Region", "Units"), ("North", 5), ("South", 7), ("East", 3)]:
    ws.append(row)
ws["A5"], ws["B5"] = "Total", "=SUM(B2:B4)"
s = wb.create_sheet("Summary")
s["A1"], s["B1"] = "South units", "=Data!B3"
s["A2"], s["B2"] = "All units", "=Data!B5"
s["A3"], s["B3"] = "South share", "=Data!B3/Data!B5"
wb.save(f"{work}/model.xlsx")
json.dump({"order": ["North", "Central", "South", "East"], "total": 19, "south": 7}, open(truth, "w"))
