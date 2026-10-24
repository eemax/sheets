import os
import openpyxl
from evalkit import Checks, args, close, values

work, truth, answer = args()
path = f"{work}/report.xlsx"
c = Checks()
if not c.add("report.xlsx exists", os.path.exists(path)):
    c.report()
    raise SystemExit
wb = openpyxl.load_workbook(path)
names = {n.lower(): n for n in wb.sheetnames}
if not c.add("Sales and Summary sheets", "sales" in names and "summary" in names, wb.sheetnames):
    c.report()
    raise SystemExit
sales, summary = wb[names["sales"]], wb[names["summary"]]


def header_row(ws):
    for row in ws.iter_rows(min_row=1, max_row=5):
        texts = [str(c.value).strip().lower() for c in row if c.value is not None]
        if "region" in texts:
            return row[0].row, {t: i + 1 for i, t in enumerate(str(c.value).strip().lower() if c.value is not None else "" for c in row)}
    return None, {}


# Sales: a Revenue formula on every data row, computing Units x Price.
h, cols = header_row(sales)
rev = cols.get("revenue")
if c.add("Sales has a Revenue column", h and rev, cols):
    data = [r for r in range(h + 1, sales.max_row + 1) if sales.cell(r, cols["region"]).value in truth["revenue"]]
    c.add("60 sales rows", len(data) == truth["rows"], len(data))
    c.add("Revenue is a formula on every row", all(isinstance(sales.cell(r, rev).value, str) and sales.cell(r, rev).value.startswith("=") for r in data))
    vals, _ = values(path, f"'{sales.title}'")
    ok = all(close(vals[r][rev - 1], vals[r][cols["units"] - 1] * vals[r][cols["price"] - 1]) for r in data)
    c.add("Revenue = Units x Price", ok)

# Summary: per region, formulas reading Sales, right values.
h, cols = header_row(summary)
need = ["region", "revenue", "target", "gap"]
if c.add("Summary has Region, Revenue, Target, Gap", h and all(k in cols for k in need), cols):
    vals, _ = values(path, f"'{summary.title}'")
    rows = {vals[r][cols["region"] - 1]: r for r in vals if r > h and vals[r] and vals[r][cols["region"] - 1] in truth["revenue"]}
    c.add("a row per region", set(rows) == set(truth["revenue"]), sorted(rows))
    for region, r in sorted(rows.items()):
        v = vals[r]
        f = summary.cell(r, cols["revenue"]).value
        c.add(f"{region}: revenue from Sales by formula", isinstance(f, str) and f.startswith("=") and "sales" in f.lower()
              and close(v[cols["revenue"] - 1], truth["revenue"][region]), (f, v[cols["revenue"] - 1]))
        g = summary.cell(r, cols["gap"]).value
        c.add(f"{region}: target and gap", close(v[cols["target"] - 1], truth["targets"][region]) and isinstance(g, str) and g.startswith("=")
              and close(v[cols["gap"] - 1], truth["revenue"][region] - truth["targets"][region]), (v[cols["target"] - 1], g, v[cols["gap"] - 1]))
    money = [summary.cell(r, cols[k]).number_format for r in rows.values() for k in ("revenue", "target", "gap")]
    c.add("Revenue, Target, Gap in dollars", money and all("$" in m for m in money), sorted(set(money)))
    gap_col = openpyxl.utils.get_column_letter(cols["gap"])
    red = False
    for cf in summary.conditional_formatting:
        if gap_col not in str(cf.sqref):
            continue
        for rule in cf.rules:
            font = rule.dxf.font if rule.dxf else None
            rgb = getattr(getattr(font, "color", None), "rgb", None) or ""
            negative = (rule.type == "cellIs" and rule.operator == "lessThan" and rule.formula == ["0"]) or \
                       (rule.type == "expression" and rule.formula and rule.formula[0].replace(" ", "").endswith("<0"))
            if negative and isinstance(rgb, str) and len(rgb) >= 6 and int(rgb[-6:-4], 16) > 150 and int(rgb[-4:-2], 16) < 100:
                red = True
    c.add("negative Gap in red (a conditional format)", red)
    c.add("Summary revenue stored", all(close(openpyxl.load_workbook(path, data_only=True)[summary.title].cell(r, cols["revenue"]).value or 0, truth["revenue"][g])
                                        for g, r in rows.items()), kind="stored")

for ws in (sales, summary):
    hr, _ = header_row(ws)
    c.add(f"{ws.title}: bold header", hr and all(cell.font.b for cell in ws[hr] if cell.value is not None))
    c.add(f"{ws.title}: header stays in view", hr and ws.freeze_panes and openpyxl.utils.cell.coordinate_from_string(ws.freeze_panes)[1] == hr + 1, ws.freeze_panes)
c.report()
