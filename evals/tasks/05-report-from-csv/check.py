import os, zipfile
import openpyxl
from evalkit import Checks, args, close, values

work, truth, answer = args()
path = f"{work}/summary.xlsx"
c = Checks()
if not c.add("summary.xlsx exists", os.path.exists(path)):
    c.report()
    raise SystemExit
wb = openpyxl.load_workbook(path)
regions = set(truth["units"])
found = None
for ws in wb.worksheets:
    rows, cols = values(path, f"'{ws.title}'")
    for r, v in rows.items():
        if v and v[0] in regions:
            found = (ws, rows)
            break
    if found:
        break
if not c.add("a sheet with the regions", found is not None):
    c.report()
    raise SystemExit
ws, rows = found
def numbers(v):
    return [x for x in v[1:] if isinstance(x, (int, float))]
for region in sorted(regions):
    v = next((v for v in rows.values() if v and v[0] == region), [])
    ns = numbers(v)
    c.add(f"{region} units and revenue", any(close(n, truth["units"][region]) for n in ns) and any(close(n, truth["revenue"][region]) for n in ns), v)
tot = next((r for r, v in rows.items() if v and isinstance(v[0], str) and v[0].lower().startswith("total")), None)
if c.add("a total row", tot is not None):
    ns = numbers(rows[tot])
    c.add("grand totals", any(close(n, sum(truth["units"].values())) for n in ns) and any(close(n, sum(truth["revenue"].values())) for n in ns), rows[tot])
header = min(rows)
c.add("bold header", all(cell.font and cell.font.b for cell in ws[header] if cell.value is not None), header)
from openpyxl.utils import column_index_from_string
_, cols = values(path, f"'{ws.title}'")
rev_cells = [ws.cell(r, column_index_from_string(cols[i])) for r, v in rows.items() if v and v[0] in regions
             for i, x in enumerate(v) if i > 0 and close(x, truth["revenue"][v[0]])]
c.add("revenue as currency", rev_cells and all(any(s in (cell.number_format or "") for s in "$€£") for cell in rev_cells),
      [cell.number_format for cell in rev_cells][:3])
with zipfile.ZipFile(path) as z:
    c.add("a chart", any(n.startswith("xl/charts/chart") for n in z.namelist()))
c.report()
