import openpyxl
from openpyxl.utils import range_boundaries
from evalkit import Checks, args, close, stored, values

work, truth, answer = args()
path = f"{work}/tables.xlsx"
c = Checks()
if not c.guard("opens", lambda: (openpyxl.load_workbook(path) is not None, "")):
    c.report()
    raise SystemExit
ws = openpyxl.load_workbook(path)["Sales"]
t = ws.tables.get("SalesT")
c.add("the table is there", t is not None)
rows, _ = values(path, "Sales")
central = [r for r, v in rows.items() if v and v[0] == "Central"]
if t is not None and c.add("Central is added once", len(central) == 1, central):
    c1, r1, c2, r2 = range_boundaries(t.ref)
    totals = t.totalsRowCount or 0
    c.add("Central is inside the table's data", r1 < central[0] <= r2 - totals, f"table {t.ref}, Central row {central[0]}")
    c.add("the totals row is kept", totals == 1 and rows.get(r2, [None])[0] == "Total", f"totalsRowCount {totals}, last row {rows.get(r2)}")
    v = rows[central[0]]
    c.add("Central's values", v[1] == 6 and close(v[2], 2), v)
    tot = next((r for r, v in rows.items() if v and v[0] == "Total"), None)
    c.add("table total", tot and close(rows[tot][1], truth["units"]), rows.get(tot))
    c.add("formula reading the table", close(rows[2][4], truth["units"]), rows[2][4])
    c.add("table total stored", tot and close(stored(path, "Sales", f"B{tot}"), truth["units"]), tot and stored(path, "Sales", f"B{tot}"), kind="stored")
c.report()
